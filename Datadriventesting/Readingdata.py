'''

Data Driven Testing is a software testing approach in  which test data is separated from test scripts, allowing the same test to be run multiple times with different sets of input data

openpyxl -> we can work with excel files

1, Read data from Excel
2, How to write data into excel
3, Data driven test case

'''

import openpyxl

#file ->  workbook -> sheet -> row -> cells
file =  r'C:\Users\ajjay\PycharmProjects\Selenium_practice\Datadriventesting\book_data.xlsx'
workbook =  openpyxl.load_workbook(file)
sheet = workbook['Sheet1']

rows = sheet.max_row
columns = sheet.max_column

for row in range(1, rows+1):
    for col in range(1, columns+1):
        print(sheet.cell(row,col).value, end='                   ')


    print()




