import openpyxl
from openpyxl.styles import PatternFill



def getRowCount(file,sheetName):
    workbook = openpyxl.load_workbook(file)
    sheet =  workbook[sheetName]
    return(sheet.max_row)

def getColumnCount(file,sheetName):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetName]
    return(sheet.max_column)

def readData(file,SheetName,rownum,colmnno):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[SheetName]
    return sheet.cell(rownum,colmnno).value

def WriteData(file, sheetname, rownum, columno, data):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    sheet.cell(rownum,columno).value = data
    workbook.save(file)

def fillGreenColour(file,sheetname,rownum,columno):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    greenfill = PatternFill(start_color='60b212',
                            end_color='60b212',
                            fill_type='solid'
                            )
    sheet.cell(rownum,columno).fill = greenfill
    workbook.save(file)


def fillRedColour(file, sheetname, rownum, columno):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    Redfill = PatternFill(start_color='ff0000',
                            end_color='ff0000',
                            fill_type='solid'
                            )
    sheet.cell(rownum, columno).fill = Redfill
    workbook.save(file)


