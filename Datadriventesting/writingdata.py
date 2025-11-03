import openpyxl

# #samedata
# file  =  r'C:\Users\ajjay\PycharmProjects\Selenium_practice\Datadriventesting\wdata1.xlsx'
# # workbook = openpyxl.load_workbook(file)
# # sheet =  workbook.active
# #
# # for r in range(1,26):
# #     for c in range(1,24):
# #         sheet.cell(r,c).value =  'Ajjayya'
# #
# # workbook.save(file)

file  =  r'C:\Users\ajjay\PycharmProjects\Selenium_practice\Datadriventesting\wdata1.xlsx'
workbook = openpyxl.load_workbook(file)
sheet =  workbook.active

sheet.cell(1,1).value = 1234
sheet.cell(1,2).value ="Ajay"
sheet.cell(1,3).value ="QA Engineer"
sheet.cell(2,1).value = 1235
sheet.cell(2,2).value="Hema"
sheet.cell(2,3).value="Mom"
sheet.append([12346,"Jeeva","Brother"])


workbook.save(file)



